from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import os
from pathlib import Path
from typing import Any
from uuid import uuid4

import requests
from fastapi import UploadFile

from app.config import ROOT_DIR
from app.llm.client import LLMClient

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {
    '.txt', '.md', '.markdown',
    '.py', '.js', '.jsx', '.ts', '.tsx',
    '.java', '.cpp', '.c', '.h', '.hpp',
    '.cs', '.go', '.rs', '.rb', '.php',
    '.swift', '.kt', '.sql',
    '.html', '.css', '.scss',
    '.json', '.yaml', '.yml', '.xml',
    '.csv', '.log',
}

HF_API_URL = "https://api-inference.huggingface.co/pipeline/feature-extraction/{model}"
DEFAULT_EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"


class RAGService:
    def __init__(self) -> None:
        self._collection_name = 'algolens_knowledge'
        self._storage_path = ROOT_DIR / 'storage' / 'chroma'
        self._embedding_model_name = DEFAULT_EMBEDDING_MODEL
        self._chunk_size = 900
        self._chunk_overlap = 160
        self._splitter = self._build_splitter()
        self._llm = LLMClient()
        self._collection = None
        self._hf_token = os.getenv('HF_TOKEN', '')
        self._graph = None

    @property
    def embedding_model_name(self) -> str:
        return self._embedding_model_name

    def _build_splitter(self):
        try:
            from langchain_text_splitters import RecursiveCharacterTextSplitter
        except ImportError:
            logger.warning(
                'LangChain text splitters are not installed; using fallback chunking.'
            )
            return None
        return RecursiveCharacterTextSplitter(
            chunk_size=self._chunk_size,
            chunk_overlap=self._chunk_overlap,
            separators=['\n\n', '\n', '. ', ' ', ''],
        )

    def _split_text(self, text: str) -> list[str]:
        if self._splitter is not None:
            return self._splitter.split_text(text)
        chunks: list[str] = []
        start = 0
        while start < len(text):
            end = start + self._chunk_size
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
            start += self._chunk_size - self._chunk_overlap
            if start <= (end - self._chunk_size):
                break
        return chunks

    def _get_collection(self):
        if self._collection is not None:
            return self._collection
        try:
            import chromadb
        except ImportError as exc:
            raise RuntimeError(
                'ChromaDB is not installed. Run pip install -r backend/requirements.txt.'
            ) from exc

        use_persistent = os.getenv('CHROMA_PERSISTENT', 'false').lower() == 'true'
        if use_persistent:
            self._storage_path.mkdir(parents=True, exist_ok=True)
            client = chromadb.PersistentClient(path=str(self._storage_path))
        else:
            client = chromadb.Client()

        self._collection = client.get_or_create_collection(
            name=self._collection_name,
            metadata={'hnsw:space': 'cosine'},
        )
        return self._collection

    def _embed(self, texts: list[str]) -> list[list[float]]:
        """Generate embeddings using the HuggingFace Inference API (free)."""
        if not self._hf_token:
            raise RuntimeError(
                'HF_TOKEN is not set. Get a free token at https://huggingface.co/settings/tokens '
                'and add it to your .env or environment variables.'
            )
        url = HF_API_URL.format(model=self._embedding_model_name)
        headers = {"Authorization": f"Bearer {self._hf_token}"}

        batch_size = 32
        all_embeddings: list[list[float]] = []
        for i in range(0, len(texts), batch_size):
            batch = texts[i : i + batch_size]
            response = requests.post(
                url,
                headers=headers,
                json={"inputs": batch, "options": {"wait_for_model": True}},
                timeout=60,
            )
            if response.status_code != 200:
                raise RuntimeError(
                    f'HuggingFace embedding API error ({response.status_code}): {response.text[:300]}'
                )
            result = response.json()
            if isinstance(result, list) and len(result) > 0:
                if isinstance(result[0], list) and isinstance(result[0][0], float):
                    all_embeddings.extend(result)
                elif isinstance(result[0], list) and isinstance(result[0][0], list):
                    all_embeddings.extend([r[0] for r in result])
                else:
                    all_embeddings.extend(result)
            else:
                raise RuntimeError(f'Unexpected HuggingFace API response format: {str(result)[:300]}')
        return all_embeddings

    async def ingest_files(self, files: list[UploadFile]) -> dict[str, Any]:
        collection = self._get_collection()
        ids: list[str] = []
        documents: list[str] = []
        metadatas: list[dict[str, Any]] = []
        skipped: list[str] = []

        for file in files:
            filename = Path(file.filename or 'uploaded-file').name
            payload = await file.read()
            try:
                text = self._extract_text(filename, payload)
            except Exception as exc:
                logger.warning('Skipping %s during ingestion: %s', filename, exc)
                skipped.append(f'{filename}: {exc}')
                continue

            clean_text = self._normalize_text(text)
            if not clean_text:
                skipped.append(f'{filename}: no extractable text')
                continue

            chunks = self._split_text(clean_text)
            source_id = self._source_id(filename, payload)
            file_type = Path(filename).suffix.lower().lstrip('.') or 'text'
            for index, chunk in enumerate(chunks):
                chunk_id = f'{source_id}:{index}:{uuid4().hex[:8]}'
                ids.append(chunk_id)
                documents.append(chunk)
                metadatas.append(
                    {
                        'source_id': source_id,
                        'filename': filename,
                        'file_type': file_type,
                        'chunk_index': index,
                    }
                )

        if documents:
            collection.add(
                ids=ids,
                documents=documents,
                metadatas=metadatas,
                embeddings=self._embed(documents),
            )

        return {
            'files_processed': len(files) - len(skipped),
            'chunks_indexed': len(documents),
            'skipped_files': skipped,
            'message': 'Knowledge base updated' if documents else 'No chunks were indexed',
        }

    def _get_graph(self):
        """Lazy-load the LangGraph workflow."""
        if self._graph is not None:
            return self._graph
        try:
            from app.services.graph_workflow import build_rag_graph
            self._graph = build_rag_graph(
                retrieve_fn=self.retrieve,
                llm=self._llm,
                top_k=6,
            )
            logger.info('LangGraph RAG workflow initialized.')
        except ImportError:
            logger.warning('LangGraph not installed; falling back to direct RAG.')
            self._graph = None
        except Exception as exc:
            logger.warning('Failed to build LangGraph workflow: %s. Falling back.', exc)
            self._graph = None
        return self._graph

    def answer_question(self, question: str, top_k: int = 5) -> dict[str, Any]:
        graph = self._get_graph()
        if graph is not None:
            try:
                result = graph.invoke({"question": question, "sources": [], "is_relevant": False, "answer": ""})
                sources = result.get("sources", [])
                answer = result.get("answer", "")
                if answer:
                    return {
                        'answer': answer,
                        'sources': sources,
                        'retrieval_count': len(sources),
                    }
            except Exception as exc:
                logger.warning('LangGraph workflow failed: %s. Falling back to direct RAG.', exc)

        # Fallback: direct retrieval + LLM
        sources = self.retrieve(question, top_k=top_k)
        if not sources:
            return {
                'answer': 'No indexed knowledge found yet. Upload notes, PDFs, editorials, or code files first.',
                'sources': [],
                'retrieval_count': 0,
            }
        answer = self._llm.answer_with_context(question, sources)
        return {
            'answer': answer,
            'sources': sources,
            'retrieval_count': len(sources),
        }

    def similar_problem(self, problem_text: str, top_k: int = 5) -> dict[str, Any]:
        return {'matches': self.retrieve(problem_text, top_k=top_k)}

    def stats(self) -> dict[str, Any]:
        return {
            'chunks': self._get_collection().count(),
            'embedding_model': self._embedding_model_name,
            'vector_store': 'ChromaDB',
        }

    def retrieve(self, query: str, top_k: int = 5) -> list[dict[str, Any]]:
        collection = self._get_collection()
        if collection.count() == 0:
            return []
        result = collection.query(
            query_embeddings=self._embed([query]),
            n_results=top_k,
            include=['documents', 'metadatas', 'distances'],
        )
        docs = result.get('documents', [[]])[0]
        metadatas = result.get('metadatas', [[]])[0]
        distances = result.get('distances', [[]])[0]
        chunks: list[dict[str, Any]] = []
        for doc, metadata, distance in zip(docs, metadatas, distances):
            score = max(0.0, min(1.0, 1.0 - float(distance)))
            chunks.append(
                {
                    'source_id': str(metadata.get('source_id', 'unknown')),
                    'filename': str(metadata.get('filename', 'unknown')),
                    'file_type': str(metadata.get('file_type', 'text')),
                    'chunk_index': int(metadata.get('chunk_index', 0)),
                    'text': doc,
                    'score': round(score, 4),
                }
            )
        return chunks

    def _extract_text(self, filename: str, payload: bytes) -> str:
        extension = Path(filename).suffix.lower()
        if extension == '.pdf':
            return self._extract_pdf(payload)
        if extension == '.docx':
            return self._extract_docx(payload)
        if extension == '.pptx':
            return self._extract_pptx(payload)
        if extension in {'.xlsx', '.xlsm'}:
            return self._extract_xlsx(payload)
        if extension == '.csv':
            return self._extract_csv(payload)
        if extension == '.json':
            return self._extract_json(payload)
        if extension in TEXT_EXTENSIONS or not extension:
            return self._decode_text(payload)
        try:
            return self._decode_text(payload)
        except UnicodeDecodeError as exc:
            raise ValueError(
                f'unsupported binary file type {extension or "without extension"}'
            ) from exc

    def _extract_pdf(self, payload: bytes) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError('PDF support requires pypdf.') from exc
        reader = PdfReader(io.BytesIO(payload))
        return '\n\n'.join(page.extract_text() or '' for page in reader.pages)

    def _extract_docx(self, payload: bytes) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError('DOCX support requires python-docx.') from exc
        document = Document(io.BytesIO(payload))
        return '\n'.join(paragraph.text for paragraph in document.paragraphs)

    def _extract_pptx(self, payload: bytes) -> str:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError('PPTX support requires python-pptx.') from exc
        deck = Presentation(io.BytesIO(payload))
        lines: list[str] = []
        for slide_index, slide in enumerate(deck.slides, start=1):
            lines.append(f'Slide {slide_index}')
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text:
                    lines.append(shape.text)
        return '\n'.join(lines)

    def _extract_xlsx(self, payload: bytes) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError('XLSX support requires openpyxl.') from exc
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        rows: list[str] = []
        for sheet in workbook.worksheets:
            rows.append(f'Sheet: {sheet.title}')
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell) for cell in row if cell is not None]
                if values:
                    rows.append(' | '.join(values))
        return '\n'.join(rows)

    def _extract_csv(self, payload: bytes) -> str:
        text = self._decode_text(payload)
        reader = csv.reader(io.StringIO(text))
        return '\n'.join(' | '.join(row) for row in reader)

    def _extract_json(self, payload: bytes) -> str:
        parsed = json.loads(self._decode_text(payload))
        return json.dumps(parsed, indent=2, ensure_ascii=True)

    def _decode_text(self, payload: bytes) -> str:
        try:
            return payload.decode('utf-8')
        except UnicodeDecodeError:
            return payload.decode('latin-1')

    def _normalize_text(self, text: str) -> str:
        return '\n'.join(line.strip() for line in text.splitlines() if line.strip())

    def _source_id(self, filename: str, payload: bytes) -> str:
        digest = hashlib.sha256(payload).hexdigest()[:16]
        safe_name = hashlib.sha1(filename.encode('utf-8')).hexdigest()[:8]
        return f'{safe_name}-{digest}'
